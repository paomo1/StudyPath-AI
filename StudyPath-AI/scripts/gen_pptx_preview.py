# -*- coding: utf-8 -*-
"""把 毕设PPT.pptx 逐页还原成 HTML 预览。

本机没有 PowerPoint / LibreOffice，改完 pptx 只能靠结构校验加这个预览页肉眼核对。
还原粒度到形状级：矩形和圆角的填充与边框、文本的段落与行内样式、图片原样内嵌。
坐标按 96 DPI 从 EMU 换算，幻灯片 13.333x7.5 英寸对应 1280x720 像素。
"""
import base64
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

PPTX = Path(r"F:\13.答辩材料\毕设PPT.pptx")
OUT = Path(r"F:\13.答辩材料\PPT版式预览.html")

DPI = 96.0
EMU_PER_INCH = 914400.0

ALIGN = {1: "center", 2: "right", 3: "justify"}
ANCHOR = {1: "flex-start", 3: "center", 4: "flex-end"}


def px(v):
    """EMU 或 Length -> CSS 像素。"""
    if v is None:
        return 0.0
    return round(v / EMU_PER_INCH * DPI, 2)


def hex_of(color_format):
    try:
        c = color_format.rgb
        return "#%02X%02X%02X" % (c[0], c[1], c[2])
    except Exception:
        return None


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def shape_fill(sh):
    try:
        t = sh.fill.type
    except Exception:
        return None
    if t is None:
        return None
    try:
        return hex_of(sh.fill.fore_color)
    except Exception:
        return None


def shape_line(sh):
    try:
        if sh.line.fill.type is None:
            return None
        return hex_of(sh.line.color)
    except Exception:
        return None


def line_width(sh):
    try:
        w = sh.line.width
        return max(0.5, round(px(w) * 0.75, 2)) if w else 0.75
    except Exception:
        return 0.75


def text_block(sh):
    tf = sh.text_frame
    anchor = ANCHOR.get(int(tf.vertical_anchor or 1), "flex-start") if tf.vertical_anchor else "flex-start"
    paras = []
    for p in tf.paragraphs:
        align = "left"
        try:
            if p.alignment is not None:
                align = ALIGN.get(int(p.alignment), "left")
        except Exception:
            pass
        spans = []
        first_size = None
        for r in p.runs:
            f = r.font
            # px() 已把 pt 换算成 CSS 像素（1pt = 4/3px），不要再乘一次
            size = px(f.size) if f.size is not None else 15.3
            if first_size is None:
                first_size = size
            color = hex_of(f.color) if f.color is not None else None
            weight = "600" if f.bold else "400"
            style = "font-size:%spx;font-weight:%s" % (round(size, 1), weight)
            if color:
                style += ";color:%s" % color
            if f.name == "Consolas":
                style += ";font-family:Consolas,monospace"
            spans.append('<span style="%s">%s</span>' % (style, esc(r.text)))
        if not spans:
            spans.append("&nbsp;")
        # 行高按本段实际字号算：挂在 <p> 上会以 body 的 16px 为基准，行距会被放大
        factor = 1.15
        try:
            if p.line_spacing and isinstance(p.line_spacing, float):
                factor = p.line_spacing
        except Exception:
            pass
        lh = "line-height:%spx" % round((first_size or 11.3) * factor, 1)
        paras.append('<p style="margin:0 0 1px 0;text-align:%s;%s">%s</p>' % (align, lh, "".join(spans)))
    return '<div class="tb" style="justify-content:%s">%s</div>' % (anchor, "".join(paras))


def render_slide(slide, idx):
    parts = []
    for sh in slide.shapes:
        st = sh.shape_type
        l, t, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if w <= 0 or h <= 0:
            continue
        base = "position:absolute;left:%spx;top:%spx;width:%spx;height:%spx;" % (l, t, w, h)

        if st == MSO_SHAPE_TYPE.PICTURE:
            b64 = base64.b64encode(sh.image.blob).decode("ascii")
            ext = (sh.image.ext or "png").lower()
            if ext == "jpg":
                ext = "jpeg"
            parts.append('<img src="data:image/%s;base64,%s" style="%sobject-fit:fill">' % (ext, b64, base))
            continue

        if st == MSO_SHAPE_TYPE.LINE:
            color = shape_line(sh) or "#999999"
            parts.append('<div style="%sborder-top:%spx solid %s"></div>' % (base, line_width(sh), color))
            continue

        if st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.FREEFORM):
            fill = shape_fill(sh)
            line = shape_line(sh)
            css = base
            if fill:
                css += "background:%s;" % fill
            if line:
                css += "border:%spx solid %s;" % (line_width(sh), line)
            name = ""
            try:
                name = sh.name or ""
            except Exception:
                pass
            if "Rounded" in name:
                css += "border-radius:6px;"
            if sh.has_text_frame and sh.text_frame.text.strip():
                parts.append('<div style="%s;overflow:hidden">%s</div>' % (css, text_block(sh)))
            else:
                parts.append('<div style="%s"></div>' % css)
            continue

        if st == MSO_SHAPE_TYPE.GROUP:
            continue

    return '<div class="page"><div class="num">%02d</div>%s</div>' % (idx, "".join(parts))


def main():
    import sys
    argv = sys.argv[1:]
    only = int(argv[argv.index("--only") + 1]) if "--only" in argv else None

    prs = Presentation(str(PPTX))
    W, H = px(prs.slide_width), px(prs.slide_height)
    pages = [render_slide(s, i) for i, s in enumerate(prs.slides, 1)]

    css = """
    body { margin:0; padding:24px 0 40px; background:#8B93A7;
           font-family:"Microsoft YaHei","Segoe UI",sans-serif; }
    .page { position:relative; width:__W__px; height:__H__px; margin:0 auto 26px;
            background:#fff; box-shadow:0 6px 26px rgba(0,0,0,.28); overflow:hidden; }
    .num { position:absolute; left:-46px; top:6px; color:#fff; font-size:13px;
           font-weight:700; opacity:.85; }
    .tb { position:absolute; inset:0; display:flex; flex-direction:column;
          padding:1px 4px; box-sizing:border-box; }
    img { image-rendering:auto; }
    """
    css = css.replace("__W__", str(int(W))).replace("__H__", str(int(H)))

    def wrap(body, extra_css=""):
        return (
            "<!DOCTYPE html>\n<html lang=\"zh-CN\">\n<head>\n<meta charset=\"utf-8\">\n"
            "<title>StudyPath 答辩 PPT 版式预览</title>\n<style>" + css + extra_css
            + "</style>\n</head>\n<body>\n" + body + "\n</body>\n</html>\n"
        )

    if only:
        # 单页导出：贴边铺满窗口，供 headless 截图逐页核对
        single_css = "body{padding:0;background:#fff}.page{margin:0;box-shadow:none}.num{display:none}"
        dst = OUT.with_name("_preview_p%02d.html" % only)
        dst.write_text(wrap(pages[only - 1], single_css), encoding="utf-8")
        print("Saved single page:", dst)
        return

    OUT.write_text(wrap("\n".join(pages)), encoding="utf-8")
    print("Saved:", OUT, "slides:", len(pages), "px:", int(W), "x", int(H))


if __name__ == "__main__":
    main()
