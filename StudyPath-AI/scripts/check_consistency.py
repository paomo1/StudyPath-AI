# -*- coding: utf-8 -*-
"""一致性校验：作废口径零命中 + 计数事实表 + 关键数字落点。

起因：P2 一处写「四个障碍」、另两处写「三个障碍」，页内自相矛盾，靠肉眼看不出来。
本脚本把「可数的事实」和「不能出现的旧口径」都写成表，机器扫一遍。

用法：
    C:\\Users\\13656\\anaconda3\\envs\\ai-base\\python.exe scripts/check_consistency.py
"""
import io
import re
import sys
from pathlib import Path

from pptx import Presentation

ANS = Path(r"F:\13.答辩材料")
PPTX = Path(sys.argv[1]) if len(sys.argv) > 1 else ANS / "毕设PPT.pptx"

TEXT_SOURCES = {
    "讲稿": ANS / "答辩演讲稿与导师问答.md",
    "核查报告": ANS / "StudyPath_全面定档核查报告.html",
    "速答卡": ANS / "答辩三问速答卡.html",
    "踩坑全记录": ANS / "StudyPath_踩坑全记录.html",
}

# ---- 1. 作废口径：只查 PPT（交付物不该出现）----
# 讲稿 / 核查报告里出现这些词是正常的 —— 它们写在「⚠️ 绝对不要说的」红线表里。
DEPRECATED_IN_PPT = [
    "3.119", "Rerank", "美英港新", "20.2 MB", "Top10", "Waitlist",
    "忠实度", "pgvector", "PostgreSQL", "37 校", "107 步", "165 条",
    "四个障碍", "格式与口径不统一",
]

# ---- 2. 计数事实表 ----
# (显示名, 匹配正则, 权威值, 是否强校验)
CHECKS = [
    ("障碍", r"障碍", 3, True),
    ("申请阶段", r"阶段", 5, True),
    ("worker", r"worker", 3, True),
    ("角色", r"角色", 3, True),
    ("知识库", r"知识库", 3, True),
    ("LangGraph 节点", r"(?<!N8N )节点", 7, False),
    ("LangGraph 边", r"边(?!界)", 8, False),
    ("能力边界", r"边界", 3, False),
]

# 红线语境：出现在这些词附近的数字是「不要这么说」的举例，不算矛盾
REDLINE_CTX = ["不要", "不能", "❌", "绝对", "作废", "红线", "旧口径", "已移除", "误写", "已改成"]

# ---- 3. 关键数字：PPT 必须出现（防整页重构时丢数）----
MUST_APPEAR = [
    "180", "144", "36", "129", "15", "20",
    "504.6", "0.066", "19.3", "51", "106.8", "1.2952", "0.7597",
    "2.98", "0.154", "0.933", "63.3", "77.6",
    "38", "22",
]

CN = {"一": 1, "两": 2, "二": 2, "三": 3, "四": 4, "五": 5,
      "六": 6, "七": 7, "八": 8, "九": 9}
QTY = "个|条|项|大|类|种|张|步|层|所|道|路|份|门|套"


def to_int(s):
    s = s.strip()
    if s.isdigit():
        return int(s)
    if s == "十":
        return 10
    if "十" in s:
        a, b = s.split("十")
        return (CN.get(a, 1) if a else 1) * 10 + (CN.get(b, 0) if b else 0)
    return CN.get(s)


def ppt_pages():
    prs = Presentation(str(PPTX))
    out = []
    for i, s in enumerate(prs.slides, 1):
        t = [sh.text_frame.text for sh in s.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
        out.append(("P%d" % i, "\n".join(t)))
    return out


def strip_html(p):
    s = io.open(p, encoding="utf-8").read()
    s = re.sub(r"<script.*?</script>", " ", s, flags=re.S)
    s = re.sub(r"<style.*?</style>", " ", s, flags=re.S)
    s = re.sub(r"<[^>]+>", " ", s)
    for a, b in [("&nbsp;", " "), ("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"')]:
        s = s.replace(a, b)
    return s


def load_sources():
    src = {}
    for name, path in TEXT_SOURCES.items():
        if not path.exists():
            continue
        src[name] = strip_html(path) if path.suffix == ".html" else io.open(path, encoding="utf-8").read()
    return src


def scan_counts(label, text, hits):
    for name, pat_str, expect, strict in CHECKS:
        pat = re.compile(r"([一二两三四五六七八九十0-9]{1,3})\s*(?:" + QTY + r")?\s*" + pat_str)
        for m in pat.finditer(text):
            v = to_int(m.group(1))
            if v is None:
                continue
            # 「同一知识库」「唯一一点」这类词里的数字不是数量断言
            prev = text[m.start() - 1] if m.start() > 0 else ""
            if len(m.group(1)) == 1 and prev in "同唯统":
                continue
            ctx = text[max(0, m.start() - 26):m.end() + 26].replace("\n", " ")
            redline = any(k in ctx for k in REDLINE_CTX)
            hits.append((label, name, v, expect, strict, redline, ctx))


def main():
    pages = ppt_pages()
    numeric = {lab: re.sub(r"[,\s]", "", t) for lab, t in pages}
    sources = load_sources()
    problems = []

    print("=" * 78)
    print("作废口径扫描（只查 PPT，应全部零命中）")
    print("=" * 78)
    for bad in DEPRECATED_IN_PPT:
        hit = [lab for lab, t in pages if bad in t]
        if hit:
            print("  [!!] %-18s -> %s" % (bad, hit))
            problems.append("PPT 作废口径残留：%s @ %s" % (bad, hit))
        else:
            print("  [ok] %-18s 零命中" % bad)

    print()
    print("=" * 78)
    print("计数事实表（同概念的数量必须等于权威值；红线语境已跳过）")
    print("=" * 78)
    hits = []
    for lab, t in pages:
        scan_counts("PPT " + lab, t, hits)
    for name, t in sources.items():
        scan_counts(name, t, hits)

    for name, pat_str, expect, strict in CHECKS:
        rows = [h for h in hits if h[1] == name]
        real = [r for r in rows if not r[5]]
        bad = [r for r in real if r[2] != expect]
        tag = "强校验" if strict else "参考"
        print("\n  -- %s（权威值 %d · %s）命中 %d 处（红线语境 %d 处已跳）"
              % (name, expect, tag, len(rows), len(rows) - len(real)))
        for lab, _, v, ex, _, redline, ctx in rows:
            mark = "  <<< 不符" if (v != ex and not redline) else ("  [红线语境]" if redline else "")
            if v != ex and not redline and not strict:
                mark = "  (参考项，不计入)"
            print("     %-12s %d%s | %s" % (lab, v, mark, ctx.strip()[:62]))
        if bad and strict:
            problems.append("计数不一致：%s -> %s" % (name, [(r[0], r[2]) for r in bad]))

    print()
    print("=" * 78)
    print("关键数字落点（PPT 应出现，防重构丢数）")
    print("=" * 78)
    for n in MUST_APPEAR:
        where = [lab for lab, t in numeric.items() if n in t]
        if where:
            print("  [ok] %-10s -> %s" % (n, ",".join(where)))
        else:
            print("  [!!] %-10s 未在 PPT 中出现" % n)
            problems.append("关键数字丢失：" + n)

    print()
    print("=" * 78)
    if problems:
        print("共 %d 处待处理：" % len(problems))
        for p in problems:
            print("  -", p)
        return 1
    print("全部通过。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
