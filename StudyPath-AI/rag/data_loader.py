# -*- coding: utf-8 -*-
"""
data_loader.py — 把 xlsx 三库读成 langchain Document 列表。

设计要点：
- 每条「记录」本身就是一个语义完整的单元（一个院校项目 / 一个录取案例 / 一篇文书），
  所以一条记录 = 一个 Document，不做切碎，检索更准。
- page_content 用自然语言把关键字段拼出来，方便 embedding 理解。
- metadata 保留结构化字段（来源 sheet、学校、项目、结果、来源 URL），
  后面可用来做按国家/学校过滤，或回答时附引用。
"""
import openpyxl
from langchain_core.documents import Document

from config import XLSX_PATH, SHEET_SCHOOLS, SHEET_CASES, SHEET_ESSAYS


# 学校中英别名映射（小写 key）。让 "CMU/帝国理工/南洋理工/港大" 等中文简称 query 能精确命中院校库
_SCHOOL_ALIAS = {
    "carnegie mellon university": "CMU（卡内基梅隆大学）",
    "massachusetts institute of technology": "MIT（麻省理工学院）",
    "stanford university": "Stanford（斯坦福大学）",
    "university of california, berkeley": "UCB（加州大学伯克利分校）",
    "university of illinois urbana-champaign": "UIUC（伊利诺伊大学香槟分校）",
    "california institute of technology": "Caltech（加州理工学院）",
    "georgia institute of technology": "Georgia Tech（佐治亚理工学院）",
    "imperial college london": "Imperial（帝国理工学院）",
    "university college london": "UCL（伦敦大学学院）",
    "university of oxford": "Oxford（牛津大学）",
    "university of cambridge": "Cambridge（剑桥大学）",
    "nanyang technological university": "NTU（南洋理工大学）",
    "national university of singapore": "NUS（新加坡国立大学）",
    "the hong kong polytechnic university": "PolyU（香港理工大学）",
    "city university of hong kong": "CityU（香港城市大学）",
    "the university of hong kong": "HKU（香港大学）",
    "university of waterloo": "Waterloo（滑铁卢大学）",
    "university of toronto": "UofT（多伦多大学）",
    "new york university": "NYU（纽约大学）",
    "cornell university": "Cornell（康奈尔大学）",
    "university of california, san diego": "UCSD（加州大学圣地亚哥分校）",
    "university of california, los angeles": "UCLA（加州大学洛杉矶分校）",
    "university of california, santa barbara": "UCSB（加州大学圣巴巴拉分校）",
    "university of washington": "UW（华盛顿大学）",
    "kaist": "KAIST（韩国科学技术院）",
    "seoul national university": "SNU（首尔国立大学）",
    "the university of tokyo": "U Tokyo（东京大学）",
    "university of queensland": "UQ（昆士兰大学）",
    "university of adelaide": "Adelaide（阿德莱德大学）",
    "eindhoven university of technology": "TU/e（埃因霍温理工大学）",
    "tsinghua university": "清华（清华大学）",
    "peking university": "北大（北京大学）",
    "university of california, davis": "UC Davis（加州大学戴维斯分校）",
    "university of california, irvine": "UCI（加州大学欧文分校）",
    "university of southern california": "USC（南加州大学）",
    "university of michigan": "UMich（密歇根大学）",
    "the university of texas at austin": "UT Austin（德克萨斯大学奥斯汀分校）",
    "university of north carolina at chapel hill": "UNC（北卡罗来纳大学教堂山分校）",
}


def _read_sheet(sheet_name):
    """读一个 sheet，返回 (header_tuple, data_rows_list)"""
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    data = rows[1:]
    return header, data


def load_schools():
    header, data = _read_sheet(SHEET_SCHOOLS)
    docs = []
    for r in data:
        if not r or not r[0]:
            continue
        d = dict(zip(header, r))
        school = str(d.get('school_name', '') or '')
        program = str(d.get('program_name', '') or '')
        # 加中英别名映射：常见学校缩写直接拼到 content 顶部，大幅提升 query "CMU/帝国理工/南大" 的召回
        alias = _SCHOOL_ALIAS.get(school.lower(), school)
        content = (
            f"{alias}（{school}） 的 {program} 项目（{d.get('degree','')}），"
            f"位于{d.get('country','')}，{d.get('rank_source','')}排名 {d.get('school_rank','')}。"
            f"所属学院：{d.get('department','')}。学制：{d.get('duration_years','')} 年。"
            f"申请截止：{d.get('deadline','')}。"
            f"最低申请要求：GPA {d.get('gpa_min','')}，托福 {d.get('toefl_min','')}，"
            f"雅思 {d.get('ielts_min','')}，GRE {d.get('gre_required','')}。"
            f"偏好背景：{d.get('background_pref','')}。"
            f"学费：约 {d.get('tuition_usd_year','')} 美元/年。"
            f"标签：{d.get('tags','')}。"
        )
        docs.append(Document(
            page_content=content,
            metadata={
                "source_sheet": "院校项目库",
                "school_name": str(d.get('school_name', '') or ''),
                "program_name": str(d.get('program_name', '') or ''),
                "country": str(d.get('country', '') or ''),
                "school_rank": str(d.get('school_rank', '') or ''),
                "source_url": str(d.get('source_url', '') or ''),
            },
        ))
    return docs


def load_cases():
    header, data = _read_sheet(SHEET_CASES)
    docs = []
    for r in data:
        if not r or not r[0]:
            continue
        d = dict(zip(header, r))
        content = (
            f"录取案例 {d.get('case_id','')}：{d.get('school_name','')} {d.get('program_name','')}。"
            f"申请者背景：{d.get('applicant_tier','')}档；GPA {d.get('gpa','')}，"
            f"托福 {d.get('toefl','')}，GRE {d.get('gre','')}，"
            f"论文 {d.get('papers','')} 篇，实习 {d.get('internships','')} 段。"
            f"录取结果：{d.get('result','')}（{d.get('year','')} 年）。"
        )
        docs.append(Document(
            page_content=content,
            metadata={
                "source_sheet": "录取案例库",
                "case_id": str(d.get('case_id', '') or ''),
                "school_name": str(d.get('school_name', '') or ''),
                "program_name": str(d.get('program_name', '') or ''),
                "result": str(d.get('result', '') or ''),
                "source_url": str(d.get('source_url', '') or ''),
            },
        ))
    return docs


def load_essays():
    header, data = _read_sheet(SHEET_ESSAYS)
    docs = []
    for r in data:
        if not r or not r[0]:
            continue
        d = dict(zip(header, r))
        content = (
            f"文书范例 {d.get('essay_id','')}：类型 {d.get('type','')}，"
            f"目标项目 {d.get('target_program','')}，质量标签 {d.get('quality_label','')}。"
            f"内容/链接：{d.get('content_or_link','')}。"
        )
        docs.append(Document(
            page_content=content,
            metadata={
                "source_sheet": "文书范例库",
                "essay_id": str(d.get('essay_id', '') or ''),
                "type": str(d.get('type', '') or ''),
                "target_program": str(d.get('target_program', '') or ''),
                "source": str(d.get('source', '') or ''),
            },
        ))
    return docs


def load_all():
    """三库合并成一个 Document 列表"""
    docs = []
    docs += load_schools()
    docs += load_cases()
    docs += load_essays()
    return docs


if __name__ == "__main__":
    all_docs = load_all()
    print(f"总文档数: {len(all_docs)}")
    for d in all_docs[:2]:
        print("---")
        print(d.page_content[:220])
        print("meta:", d.metadata)
